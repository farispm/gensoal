import io
import re
import streamlit as st
import google.generativeai as genai
from docx import Document
from docx.shared import Inches, Pt

# Konfigurasi Halaman Web
st.set_page_config(page_title="Generator Soal SD/MI AI", page_icon="📝", layout="wide")

st.title("📝 Generator Soal SD/MI Berbasis AI Generatif")
st.caption("Aplikasi pembuat naskah soal dinamis Kurikulum Merdeka & K13 untuk SD dan MI")

# Fungsi Pembersih Format Markdown Inline untuk File Word
def bersihkan_markdown_inline(teks):
    # Menghapus simbol bold (**) dan italic (*) dari teks agar rapi di Word
    teks_bersih = re.sub(r'\*\*(.*?)\*\*', r'\1', teks)
    teks_bersih = re.sub(r'\*(.*?)\*', r'\1', teks_bersih)
    return teks_bersih

# Fungsi Konversi Teks Markdown ke Dokumen Word (.docx)
def konversi_ke_docx(teks_md, judul="Naskah Soal"):
    doc = Document()
    
    # Atur Margin Halaman (1 Inci / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = teks_md.split('\n')
    in_table = False
    table_data = []

    for line in lines:
        stripped = line.strip()

        # Deteksi Baris Tabel Markdown (| ... |)
        if stripped.startswith('|') and '|' in stripped[1:]:
            if '---' in stripped:  # Lewati baris pembatas tabel
                continue
            cells = [bersihkan_markdown_inline(c.strip()) for c in stripped.split('|')[1:-1]]
            if cells:
                table_data.append(cells)
                in_table = True
            continue
        else:
            # Jika keluar dari area tabel, cetak tabel ke dokumen Word
            if in_table and table_data:
                cols = max(len(r) for r in table_data)
                table = doc.add_table(rows=len(table_data), cols=cols)
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        if c_idx < cols:
                            table.cell(r_idx, c_idx).text = val
                table_data = []
                in_table = False

        # Bersihkan format inline untuk teks paragraf/judul
        teks_paragraf = bersihkan_markdown_inline(stripped)

        # Deteksi Judul & Heading
        if stripped.startswith('# '):
            doc.add_heading(teks_paragraf[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(teks_paragraf[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(teks_paragraf[4:], level=3)
        elif stripped.startswith('#### '):
            doc.add_heading(teks_paragraf[5:], level=4)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            doc.add_paragraph(teks_paragraf[2:], style='List Bullet')
        elif stripped == '':
            continue
        else:
            doc.add_paragraph(teks_paragraf)

    # Cetak tabel jika berada di bagian akhir dokumen
    if in_table and table_data:
        cols = max(len(r) for r in table_data)
        table = doc.add_table(rows=len(table_data), cols=cols)
        table.style = 'Table Grid'
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                if c_idx < cols:
                    table.cell(r_idx, c_idx).text = val

    # Simpan ke byte stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# Panel Samping - Konfigurasi Utama
with st.sidebar:
    st.header("⚙️ Konfigurasi Utama")
    
    # Memeriksa API key dari st.secrets jika ada, atau isi kunci default Anda di sini
    default_key = st.secrets.get("GEMINI_API_KEY", "")
    
    api_key = st.text_input(
        "Gemini API Key", 
        value=default_key, 
        type="password", 
        help="Masukkan API Key Google Gemini Anda"
    )
    
    kelas = st.selectbox("Tingkat Kelas", ["Kelas I", "Kelas II", "Kelas III", "Kelas IV", "Kelas V", "Kelas VI"])
    semester = st.radio("Semester", ["Ganjil", "Genap"], horizontal=True)
    
    mapel = st.selectbox("Mata Pelajaran", [
        "Al-Quran Hadis", "Akidah Akhlak", "Fikih", "Sejarah Kebudayaan Islam (SKI)", 
        "Bahasa Arab", "Pendidikan Pancasila", "Bahasa Indonesia", "Matematika", 
        "IPAS", "PJOK", "Seni dan Budaya", "Bahasa Inggris"
    ])
    
    level_kognitif = st.multiselect(
        "Tingkat Kesulitan (Dapat Pilih Lebih dari 1)",
        ["L1 (Pemahaman/Pengetahuan)", "L2 (Aplikasi/Penerapan)", "L3 (Penalaran/HOTS)"],
        default=["L1 (Pemahaman/Pengetahuan)", "L2 (Aplikasi/Penerapan)", "L3 (Penalaran/HOTS)"]
    )
    
    opsi_gambar = st.checkbox("Sertakan Soal Berbasis Gambar / Ilustrasi", value=False)

# Form Input Utama
st.subheader("📌 Materi & Capaian Pembelajaran (Opsional)")
col_m1, col_m2 = st.columns(2)
with col_m1:
    materi = st.text_area("Materi / Bab Pembahasan", placeholder="Contoh: Bab 3 - Daur Hidup Makhluk Hidup")
with col_m2:
    cp = st.text_area("Capaian Pembelajaran (CP) / Tujuan Pembelajaran (TP)", placeholder="Contoh: Peserta didik dapat menganalisis siklus hidup hewan...")

st.subheader("📊 Konfigurasi Bentuk & Jumlah Soal")
col1, col2, col3 = st.columns(3)

with col1:
    st.markdown("**1. Pilihan Ganda (PG)**")
    jml_pg = st.number_input("Jumlah Soal PG", min_value=0, max_value=50, value=10)
    opsi_pg = st.selectbox("Banyaknya Opsi Jawaban", [3, 4, 5], index=1 if kelas in ["Kelas IV", "Kelas V", "Kelas VI"] else 0)

    st.markdown("**2. Pilihan Ganda Kompleks (PGK)**")
    jml_pgk = st.number_input("Jumlah Soal PGK", min_value=0, max_value=30, value=2)
    pernyataan_pgk = st.number_input("Jumlah Opsi Pernyataan (PGK)", min_value=3, max_value=6, value=4)

with col2:
    st.markdown("**3. Benar / Salah (BS)**")
    jml_bs = st.number_input("Jumlah Soal B/S", min_value=0, max_value=30, value=2)
    pernyataan_bs = st.number_input("Jumlah Pernyataan per Soal B/S", min_value=2, max_value=6, value=3)

    st.markdown("**4. Menjodohkan**")
    jml_jodoh = st.number_input("Jumlah Soal Menjodohkan", min_value=0, max_value=20, value=2)
    pasangan_jodoh = st.number_input("Jumlah Pasangan per Soal", min_value=3, max_value=10, value=4)

with col3:
    st.markdown("**5. Isian Singkat**")
    jml_isian = st.number_input("Jumlah Soal Isian", min_value=0, max_value=30, value=5)

    st.markdown("**6. Uraian / Essay**")
    jml_uraian = st.number_input("Jumlah Soal Uraian", min_value=0, max_value=20, value=3)

st.markdown("---")

# Eksekusi Pembuatan Soal AI
if st.button("🚀 Buat Naskah Soal & Kunci Jawaban", type="primary", use_container_width=True):
    if not api_key:
        st.error("Silakan masukkan API Key Google Gemini pada panel sebelah kiri.")
    elif not level_kognitif:
        st.error("Pilih minimal satu tingkat kesulitan (Level Kognitif).")
    else:
        try:
            # Konfigurasi API dengan Model Resmi Gemini 1.5 Flash
            genai.configure(api_key=api_key)
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt_system = f"""
Anda adalah pakar pembuat soal asesmen pendidikan dasar SD dan MI yang berpengalaman.
Tugas Anda adalah menyusun naskah soal asesmen yang autentik, variatif, dan kontekstual.
DILARANG menggunakan pola kaku atau template berulang.

SPESIFIKASI ASESMEN:
- Tingkat: {kelas} | Semester: {semester}
- Mata Pelajaran: {mapel}
- Materi Spesifik: {materi if materi else 'Menyesuaikan Kurikulum Merdeka/K13 standar untuk mapel dan kelas ini'}
- Capaian Pembelajaran (CP): {cp if cp else 'Menyesuaikan indikator CP standar nasional'}
- Level Kognitif Terintegrasi: {', '.join(level_kognitif)}
- Penggunaan Gambar: {'Aktif. Untuk soal yang memerlukan visual, sertakan petunjuk [DESKRIPSI GAMBAR: deskripsi detail objek visual yang harus ditampilkan guru]' if opsi_gambar else 'Tidak menggunakan gambar.'}

KOMPOSISI SOAL:
1. Pilihan Ganda: {jml_pg} soal, tiap soal memiliki {opsi_pg} opsi jawaban (A, B, C, dst.).
2. Pilihan Ganda Kompleks: {jml_pgk} soal, tiap soal memiliki {pernyataan_pgk} pernyataan (jawaban benar lebih dari satu).
3. Benar / Salah: {jml_bs} soal, tiap nomor terdiri dari {pernyataan_bs} pernyataan untuk dinilai Benar/Salah.
4. Menjodohkan: {jml_jodoh} paket soal, tiap paket memiliki {pasangan_jodoh} pernyataan lajur kiri dan {pasangan_jodoh} pilihan lajur kanan.
5. Isian Singkat: {jml_isian} soal.
6. Uraian / Essay: {jml_uraian} soal.

STRUKTUR KELUARAN (MARKDOWN):
1. Header Lembar Soal & Petunjuk Pengerjaan
2. Bagian A: Pilihan Ganda
3. Bagian B: Pilihan Ganda Kompleks
4. Bagian C: Benar / Salah
5. Bagian D: Menjodohkan (Format Tabel)
6. Bagian E: Isian Singkat
7. Bagian F: Uraian
8. Kunci Jawaban & Pedoman Penskoran/Rubrik
9. Tabel Kisi-Kisi Soal (Nomor Soal, CP/Materi, Indikator Soal, Bentuk Soal, Level Kognitif L1/L2/L3).
"""
            with st.spinner("AI sedang mendesain naskah soal variatif & kisi-kisi... Mohon tunggu..."):
                response = model.generate_content(prompt_system)
                
            st.success("✨ Naskah Soal Berhasil Dibuat!")
            st.markdown(response.text)
            
            # Membuat File Word (.docx)
            file_docx = konversi_ke_docx(response.text, judul=f"Soal_{mapel}_{kelas}")
            
            # Tombol Download
            col_dl1, col_dl2 = st.columns(2)
            with col_dl1:
                st.download_button(
                    label="📄 Download Naskah Word (.docx)",
                    data=file_docx,
                    file_name=f"Soal_{mapel.replace(' ', '_')}_{kelas.replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                    use_container_width=True
                )
            with col_dl2:
                st.download_button(
                    label="📥 Download Naskah Markdown (.md)",
                    data=response.text,
                    file_name=f"Soal_{mapel.replace(' ', '_')}_{kelas.replace(' ', '_')}.md",
                    mime="text/markdown",
                    use_container_width=True
                )
                
        except Exception as err:
            st.error(f"Terjadi kesalahan pada sistem AI: {str(err)}")
